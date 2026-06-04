import fitz  # PyMuPDF
import docx
import openpyxl
import io
import os
from PIL import Image
from typing import List, Dict, Any, Tuple
from config import get_settings

settings = get_settings()

# ── File Type Detector ───────────────────────────────────
def detect_file_type(filename: str, content: bytes) -> str:
    ext = filename.lower().split(".")[-1]
    mapping = {
        "pdf": "pdf",
        "png": "image", "jpg": "image", "jpeg": "image",
        "webp": "image", "bmp": "image", "tiff": "image",
        "docx": "word", "doc": "word",
        "xlsx": "excel", "xls": "excel", "csv": "excel",
        "mp3": "audio", "wav": "audio", "m4a": "audio",
        "ogg": "audio", "flac": "audio",
        "txt": "text", "md": "text", "json": "text",
        "py": "text", "js": "text", "ts": "text",
        "html": "text", "xml": "text", "csv": "text"
    }
    return mapping.get(ext, "unknown")

# ── PDF Parser ───────────────────────────────────────────
def parse_pdf(content: bytes) -> Dict[str, Any]:
    doc = fitz.open(stream=content, filetype="pdf")
    pages = []
    images = []
    total_text = ""

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text").strip()
        pages.append(text)
        total_text += text + "\n"

        # Extract images from page
        img_list = page.get_images()
        for img_index, img in enumerate(img_list):
            xref = img[0]
            base_image = doc.extract_image(xref)
            images.append({
                "page": page_num + 1,
                "index": img_index,
                "data": base_image["image"],
                "ext": base_image["ext"]
            })

    doc.close()

    is_scanned = len(total_text.strip()) < 100 and len(images) > 0

    return {
        "type": "pdf",
        "total_pages": len(pages),
        "pages": pages,
        "images": images,
        "total_text": total_text,
        "is_scanned": is_scanned,
        "needs_ocr": is_scanned
    }

# ── Word Parser ──────────────────────────────────────────
def parse_word(content: bytes) -> Dict[str, Any]:
    doc = docx.Document(io.BytesIO(content))
    paragraphs = []
    tables = []

    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            table_data.append(row_data)
        tables.append(table_data)

    full_text = "\n".join(paragraphs)

    return {
        "type": "word",
        "paragraphs": paragraphs,
        "tables": tables,
        "total_text": full_text,
        "total_pages": len(paragraphs) // 30 + 1
    }

# ── Excel / CSV Parser ───────────────────────────────────
def parse_excel(content: bytes, filename: str) -> Dict[str, Any]:
    sheets = {}

    if filename.endswith(".csv"):
        import csv
        text = content.decode("utf-8", errors="ignore")
        reader = csv.DictReader(io.StringIO(text))
        rows = list(reader)
        sheets["Sheet1"] = rows
    else:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            headers = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    headers = [str(c) if c else f"col_{j}" 
                               for j, c in enumerate(row)]
                else:
                    if any(c is not None for c in row):
                        row_dict = dict(zip(headers, row))
                        rows.append(row_dict)
            sheets[sheet_name] = rows

    # Convert to text for LLM
    text_parts = []
    for sheet, rows in sheets.items():
        text_parts.append(f"Sheet: {sheet}")
        for row in rows[:100]:  # Limit rows
            text_parts.append(str(row))

    return {
        "type": "excel",
        "sheets": sheets,
        "total_text": "\n".join(text_parts),
        "total_rows": sum(len(r) for r in sheets.values())
    }

# ── Image Parser ─────────────────────────────────────────
def parse_image(content: bytes) -> Dict[str, Any]:
    import base64
    img = Image.open(io.BytesIO(content))
    width, height = img.size
    img_b64 = base64.b64encode(content).decode()

    return {
        "type": "image",
        "width": width,
        "height": height,
        "image_base64": img_b64,
        "total_text": "",
        "needs_vision": True
    }

# ── Text Parser ──────────────────────────────────────────
def parse_text(content: bytes) -> Dict[str, Any]:
    text = content.decode("utf-8", errors="ignore")
    return {
        "type": "text",
        "total_text": text,
        "total_pages": 1
    }

# ── Audio Parser ─────────────────────────────────────────
def parse_audio(content: bytes, filename: str) -> Dict[str, Any]:
    # Save temp file for whisper
    temp_path = f"/tmp/{filename}"
    with open(temp_path, "wb") as f:
        f.write(content)
    return {
        "type": "audio",
        "temp_path": temp_path,
        "total_text": "",
        "needs_transcription": True
    }

# ── Main Parser ──────────────────────────────────────────
def parse_file(content: bytes, filename: str) -> Dict[str, Any]:
    file_type = detect_file_type(filename, content)

    if file_type == "pdf":
        return parse_pdf(content)
    elif file_type == "image":
        return parse_image(content)
    elif file_type == "word":
        return parse_word(content)
    elif file_type == "excel":
        return parse_excel(content, filename)
    elif file_type == "audio":
        return parse_audio(content, filename)
    elif file_type == "text":
        return parse_text(content)
    else:
        # Try as text
        try:
            return parse_text(content)
        except:
            return {"type": "unknown", "total_text": "", "error": "Unsupported file type"}
